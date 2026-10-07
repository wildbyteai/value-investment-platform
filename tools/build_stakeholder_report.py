"""Local report authoring. Synthetic arithmetic only; no app import, database or network."""
from pathlib import Path
from datetime import datetime
from zoneinfo import ZoneInfo
from decimal import Decimal, localcontext
import json, hashlib, subprocess, re, math
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT=Path(__file__).resolve().parents[1]
VERSION='0.6'
REPORT_DATE='2026-10-07'
SOURCES=['PROJECT.md','README.md','CONTEXT.md','docs/02-business-design.md','docs/11-templates-automation-and-roles.md','docs/12-time-numerics-and-corrections.md','docs/14-domain-and-business-flows.md','versions/v0.0.1/VERIFICATION.md','review/FUTU-INTEGRATION.md','review/REAL-GAP-CLOSURE.md','review/REAL-CLOSURE-SUPPLEMENT.md','docs/08-security-and-operations.md','research/current-share-coverage.md','config/scoring-standard-v1.json','config/strategy-standard-v1.json','reports/stakeholder/business-narrative.md','reports/stakeholder/case-study.json','reports/stakeholder/frameworks-and-costs.md','research/report-component-cost-basis.md','reports/stakeholder/screenshots-v0.6.json','docs/adr/0004-elasticsearch-search-platform.md','docs/03-architecture.md','docs/07-ai-and-retrieval.md','frontend/src/app.tsx','frontend/src/ui.tsx','frontend/src/judgments.tsx','frontend/src/template.tsx']
REV=subprocess.check_output(['git','rev-parse','HEAD'],cwd=ROOT,text=True).strip()
HASHES={p:hashlib.sha256((ROOT/p).read_bytes()).hexdigest() for p in SOURCES}
ASSETS=ROOT/'reports/stakeholder/figures/v0.6';ASSETS.mkdir(parents=True,exist_ok=True)
OUT=ROOT/'output/documents';OUT.mkdir(parents=True,exist_ok=True)
FONT='/System/Library/Fonts/Hiragino Sans GB.ttc';FONTNAME='Hiragino Sans GB'
INK='#172F42';BLUE='#185774';TEAL='#16756E';GREY='#677A89';RED='#A33C43';AMBER='#95691B';PALE='#F0F5F7'
figures=[]
def font(n):return ImageFont.truetype(FONT,n)
def lines(s,f,w,d):
    result=[]
    for raw in s.split('\n'):
        line=''
        for ch in raw:
            if d.textlength(line+ch,font=f)>w and line:result.append(line);line=ch
            else:line+=ch
        result.append(line)
    return result
def txt(d,xy,s,size=46,color=INK,width=None):
    f=font(size);x,y=xy
    for ln in lines(s,f,width or (2180-x),d):d.text((x,y),ln,font=f,fill=color);y+=size*1.35
    return y
def canvas(h):
    im=Image.new('RGB',(2200,h),'white');return im,ImageDraw.Draw(im)
def arrow(d,a,b,color=GREY,width=5):
    d.line([a,b],fill=color,width=width);ang=math.atan2(b[1]-a[1],b[0]-a[0]);d.polygon([b,(b[0]-20*math.cos(ang-.4),b[1]-20*math.sin(ang-.4)),(b[0]-20*math.cos(ang+.4),b[1]-20*math.sin(ang+.4))],fill=color)
def save(fid,im,caption):
    p=ASSETS/(fid+'.png');im.save(p,dpi=(300,300));figures.append({'id':fid,'caption':caption,'file':str(p.relative_to(ROOT)),'pixel_size':list(im.size),'kind':'local_business_infographic_or_schematic','actual_screenshot':False})

D=Decimal;case=json.loads((ROOT/'reports/stakeholder/case-study.json').read_text());scoring=json.loads((ROOT/'config/scoring-standard-v1.json').read_text());strategy=json.loads((ROOT/'config/strategy-standard-v1.json').read_text())
def linear(v,zero,full):return max(D(0),min(D(100),(v-zero)/(full-zero)*100))
EPS=D(case['ordinary_profit_ttm_cny'])/D(case['same_rights_ordinary_shares']);FX=D(case['fx_cny_per_hkd'])
base={}
for dim in ['profit_quality','financial_resilience']:
    base[dim]=sum(D(m['weight'])*linear(D(case[m['key']]),D(m['zero_at']),D(m['full_at'])) for m in scoring['baselines'][dim]['metrics'])
base.update({k:D(v) for k,v in case['rubric_scores'].items()})
Q=sum(base[k]*D(w) for k,w in scoring['dimension_weights'].items())
ef=case['event_factors'];EVENT=D(scoring['event_contribution']['scale'])*D(ef['impact'])*D(ef['relevance'])*D(ef['confidence'])*D(ef['source_quality'])
def pe(p):return D(p)*FX/EPS
def value(p):return linear(pe(p),D(scoring['valuation']['zero_score_at']),D(scoring['valuation']['full_score_at']))
def fmt(v,n=2):return f'{v:.{n}f}'
def q_at(i):
    if i==0:return Q
    if i>=6:return Q+EVENT*D(scoring['dimension_weights']['growth_sustainability'])
    with localcontext() as ctx:
        ctx.prec=28
        decay=D(2)**(-D(case['sessions'][i]['event_age_calendar_days'])/D(scoring['event_contribution']['half_life_calendar_days']['growth_sustainability']))
        return Q+EVENT*decay*D(scoring['dimension_weights']['growth_sustainability'])
assert EPS==1 and Q==79 and EVENT==4 and base['profit_quality']==92 and base['financial_resilience']==85
assert strategy['state_policy']['enter_distinct_sessions']==2 and strategy['state_policy']['unknown_action']=='hold_last_confirmed_reset_pending'
assert 'confirmed_debt_default' in strategy['hard_risk_policy']['risk_codes']

# B-01: clear handoffs, with the detailed processing contract in the adjoining table.
im,d=canvas(440)
stages=[('01','取得资料','原文 / 财务 / 行情'),('02','形成判断','事实 / 影响\n待补依据'),('03','分别评价','公司质量\n证券估值'),('04','策略跟踪','条件 / 确认 / 风险'),('05','复核历史','变化解释\n原始依据')]
for i,(num,title,output) in enumerate(stages):
    x=30+i*438;d.ellipse((x+5,30,x+100,125),fill=BLUE);txt(d,(x+23,46),num,48,'white')
    txt(d,(x,160),title,54);d.line((x,245,x+350,245),fill='#A7BCC8',width=3);txt(d,(x,278),output,43,GREY,width=355)
    if i<4:arrow(d,(x+110,78),(x+380,78),BLUE)
save('B-01',im,'输入与交接产物')

# B-02: a numerical valuation scale, rather than a four-box picture.
im,d=canvas(540);txt(d,(30,15),'统一每股盈利1元人民币；A/H经营质量共用',52)
x0,x1=180,2070;y=330
xp=lambda x:x0+float((D(str(x))-10)/15)*(x1-x0)
for a,b,col,label in [(10,16,'#DDEEEA','达到进入门'),(16,17.5,'#F5ECD8','仅保持'),(17.5,25,'#F2E2E2','未达保持门')]:
    d.rectangle((xp(a),y-52,xp(b),y+35),fill=col)
    txt(d,(xp(a)+15,y+92),label,40,TEAL if a==10 else (AMBER if a==16 else RED),width=xp(b)-xp(a)-25)
d.line((x0,y,x1,y),fill=INK,width=4)
for n in [10,16,17.5,20,25]:
    x=xp(n);d.line((x,y+35,x,y+55),fill=INK,width=3);txt(d,(x-25,y+52),str(n),37)
for p,title,col,yy in [(D('14.4'),'H股：PE14.40 / 估值分70.67',TEAL,125),(D('20'),'A股：PE20.00 / 估值分33.33',RED,205)]:
    x=xp(p);d.ellipse((x-13,y-13,x+13,y+13),fill=col);d.line((x,y-13,x,yy+60),fill=col,width=4);txt(d,(max(30,x-300),yy),title,45,col)
save('B-02',im,'A/H估值分歧的数值解释')

# B-03: evidence and fact identity are concrete fields, not just abstract nodes.
im,d=canvas(610)
columns=[(30,'公告与报道',['订单20亿元','上年收入100亿元','规模20% → 继续读履约条款']),
         (770,'同一事实记录',['远景制造＋订单签订','合同批次＋事实期间','5份出处 → 1个有效贡献']),
         (1510,'有效经营判断',['成长维度当前贡献＋4','权重10% → 总分＋0.4','79.0 → 79.4'])]
for x,title,items in columns:
    txt(d,(x,20),title,55,BLUE);d.line((x,105,x+650,105),fill=BLUE,width=6)
    yy=140
    for t in items:yy=txt(d,(x,yy),t,45,width=650)+25
arrow(d,(685,295),(750,295));arrow(d,(1425,295),(1490,295))
d.line((30,490,2150,490),fill='#BACBD3',width=3);txt(d,(30,522),'吸收前：基准60＋事件4＝64   →   吸收后：基准64＋事件0＝64（同一比较时点）',45,TEAL)
save('B-03',im,'证据、事实归并和贡献吸收')

# B-04: missing input is visibly a break in the confirmation sequence.
im,d=canvas(510)
for i,s in enumerate(case['sessions']):
    x=130+i*274;col=RED if i==7 else (AMBER if i==2 else (TEAL if i>=4 else BLUE))
    txt(d,(x-35,20),s['id'],52,col);d.ellipse((x-22,130,x+22,174),fill=col)
    txt(d,(x-110,230),s['display'],43,col,width=240);txt(d,(x-110,305),'不等普通\n确认' if i==7 else s['counter'],42,GREY,width=240)
    if i<7:
        if i in [1,2]:
            for xx in range(x+30,x+236,26):d.line((xx,152,xx+13,152),fill=AMBER,width=4)
        else:arrow(d,(x+30,152),(x+235,152),GREY)
txt(d,(35,415),'D2清零 → D3重新1/2 → D4确认2/2；D7已确认风险走独立优先处理',45,INK)
save('B-04',im,'H股确认与风险时间线')

# B-05: a compact information schematic based on source UI areas, not a fake screenshot.
im,d=canvas(1150)
def panel(x,y,w,h,title):
    d.rectangle((x,y,x+w,y+h),outline='#8BA3B0',width=3);d.rectangle((x,y,x+w,y+95),fill=BLUE);txt(d,(x+25,y+20),title,49,'white')
panel(20,20,1370,650,'公司研究  /  远景制造  /  演示D1')
txt(d,(50,140),'经营质量79.4   覆盖100%   [回查依据]',47,TEAL)
txt(d,(50,220),'经营质量  |  财务估值  |  资料研判  |  评分规则',42)
d.line((45,300,1360,300),fill='#B5C6CE',width=3)
for yy,a,b,c in [(335,'证券','PE / 估值分','本次显示 / 已确认'),(420,'SYN-A 20.00元','20.00 / 33.33','区间外 / 区间外'),(510,'SYN-H 16.00港元','14.40 / 70.67','待入选1/2 / 区间外')]:
    txt(d,(50,yy),a,42);txt(d,(570,yy),b,42);txt(d,(965,yy),c,40,width=400)
panel(1430,20,750,650,'原文与判断')
txt(d,(1460,150),'订单公告\n金额：20亿元\n上年收入：100亿元\n5篇报道 → 1个事实\n成长当前贡献：＋4\n[原文] [理由与适用期]',43,width=680)
panel(20,720,2160,395,'研究快照比较  /  D0 → D1')
txt(d,(55,850),'新增：订单事实1项\n公司质量：79.0 → 79.4',46,width=670)
txt(d,(820,850),'H价：20.00 → 16.00港元\n估值分：46.67 → 70.67',46,width=670)
txt(d,(1570,850),'状态：区间外 → 待入选\n确认：0/2 → 1/2',44,width=560)
save('B-05',im,'页面信息结构示意（合成数据，非运行截图）')

im,d=canvas(450)
for i,(title,desc) in enumerate([('授权成员','企业登录 / 角色'),('域名入口','DNS / HTTPS'),('应用与任务','页面 / API / 后台'),('研究数据','数据库 / 原件'),('恢复与监控','异机备份 / 告警')]):
    x=30+i*435;d.rectangle((x,30,x+380,105),fill=BLUE if i<3 else TEAL);txt(d,(x+15,40),title,48,'white');txt(d,(x,165),desc,43,width=380)
    if i<4:arrow(d,(x+385,65),(x+423,65))
txt(d,(30,315),'对外开放研究入口；数据库与管理端保留私网；重建应用不丢原件与历史',45)
save('B-06',im,'生产访问与恢复职责')

doc=Document();sec=doc.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=Cm(1.65);sec.bottom_margin=Cm(1.65);sec.left_margin=Cm(2);sec.right_margin=Cm(2)
sec.header_distance=Cm(.6);sec.footer_distance=Cm(.65)
for name,size in [('Normal',11.5),('Title',27),('Subtitle',14),('Heading 1',18),('Heading 2',13),('Caption',9.5)]:
    st=doc.styles[name];st.font.name=FONTNAME;st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
    st._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),FONTNAME)
    st.paragraph_format.line_spacing=1.12;st.paragraph_format.space_after=Pt(5)
    st.paragraph_format.page_break_before=False
for name in ['Heading 1','Heading 2']:
    doc.styles[name].paragraph_format.space_before=Pt(12);doc.styles[name].paragraph_format.keep_with_next=True
doc.styles['Normal'].font.bold=False
doc.styles['Caption'].font.bold=False
doc.styles['Caption'].font.italic=False
doc.styles['Caption'].paragraph_format.keep_with_next=True
for e in doc.styles._element.xpath('.//w:rFonts'):
    for a in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme','csTheme']:e.attrib.pop(qn('w:'+a),None)
    for a in ['ascii','hAnsi','eastAsia']:e.set(qn('w:'+a),FONTNAME)
for e in doc.styles._element.xpath('.//w:pBdr'):e.getparent().remove(e)
doc.core_properties.title='价值投资研究系统：业务方案与使用流程';doc.core_properties.author='项目团队'
doc.core_properties.subject='需求方审阅：业务数据流、数值案例、功能指引与运行准备'
hp=sec.header.paragraphs[0];hp.text='价值投资研究系统  |  业务方案与使用流程'
for r in hp.runs:r.font.size=Pt(9)
fp=sec.footer.paragraphs[0];fp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
fp.add_run(f'报告{VERSION}  |  {REPORT_DATE}  |  ');fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');fp._p.append(fld)
for r in fp.runs:r.font.size=Pt(9)
def para(s,style=None):return doc.add_paragraph(s,style=style)
def heading(s,level=1):return doc.add_heading(s,level=level)
def picture(fid,cap,width=17):
    p=para('');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True;p.paragraph_format.space_after=Pt(3)
    sh=p.add_run().add_picture(str(ASSETS/Path(fid).name),width=Cm(width));sh._inline.docPr.set('descr',cap)
    p=para(cap,'Caption');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
def table(headers,rows,widths=None,fs=10.5):
    n=len(headers);widths=widths or ([3.0,7.3,6.7] if n==3 else [17/n]*n)
    t=doc.add_table(rows=1,cols=n);t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for c,w in zip(t.columns,widths):c.width=Cm(w)
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    t._tbl.tblPr.append(borders)
    for c,v in zip(t.rows[0].cells,headers):c.text=v
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        for c,v in zip(t.add_row().cells,row):c.text=str(v)
    for i,row in enumerate(t.rows):
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for j,c in enumerate(row.cells):
            c.width=Cm(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp=c._tc.get_or_add_tcPr();e=OxmlElement('w:shd');e.set(qn('w:fill'),'185774' if i==0 else ('FFFFFF' if i%2 else 'F0F5F7'));cp.append(e)
            mar=OxmlElement('w:tcMar')
            for side,v in [('top','95'),('bottom','95'),('left','100'),('right','100')]:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),v);e.set(qn('w:type'),'dxa');mar.append(e)
            cp.append(mar)
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(1);p.paragraph_format.space_before=Pt(1);p.paragraph_format.line_spacing=1.08
                if i==0 or (headers[0] in ['经营维度','工作角色','日'] and i<len(t.rows)-1):p.paragraph_format.keep_with_next=True
                if widths[j]<=1.6:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:r.font.size=Pt(fs);r.font.bold=i==0;r.font.color.rgb=RGBColor.from_string('FFFFFF' if i==0 else '000000')
    p=para('');p.paragraph_format.space_after=Pt(1);p.paragraph_format.line_spacing=.3;p.paragraph_format.space_before=Pt(0)
    return t
def link(p,label,url):
    h=OxmlElement('w:hyperlink');h.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True));r=OxmlElement('w:r');pr=OxmlElement('w:rPr');co=OxmlElement('w:color');co.set(qn('w:val'),'185774');pr.append(co);r.append(pr);tx=OxmlElement('w:t');tx.text=label;r.append(tx);h.append(r);p._p.append(h)

para('价值投资研究系统','Title');para('业务方案、数值推演与使用流程','Subtitle')
para(f'需求方审阅 · 报告{VERSION} · {REPORT_DATE}','Caption')
heading('业务定位与管理层摘要',2)
para('本系统面向基本面投资团队，建设可追溯的数字证据链与研究档案。它将资料、经营判断、证券估值和策略变化连成组织研究资产，使人员交接与时间推移后仍能复核原有逻辑。公司经营与A/H证券分别评价，帮助团队解释同一企业在不同价格和币种下的配置条件。新资料、人工修订和规则变化均保留依据与版本，让研究讨论从结论争议回到事实、口径及变化原因。')
table(['需求方的问题','系统交付的答案'],[
 ('为什么要关注这家公司？','经营质量、覆盖度和支撑判断的财报/经营证据。'),
 ('为什么A股与H股结论不同？','同一经营依据下，分别核对价格、币种、每股盈利和估值条件。'),
 ('哪些变化需要继续处理？','新资料、价格、判断、规则及风险的差异，附下一步与历史依据。')],[5.0,12.0],fs=11)
heading('本报告怎样说明方案',2)
para('先讲输入、处理和输出，再以“远景制造”演示具体数字与八个时点的状态变化，随后对应主要页面、职责及交付路径。最后用附件说明生产资源、框架与费用，并以六张实际页面截图对应日常工作。')
para('当前为开发与业务联调阶段，进展及交付条件集中在第8节。首版交付研究与状态变化记录，通知后置，不产生交易指令。')
heading('沙盘口径',2)
para('远景制造及其证券、金额、价格和判断参数均为合成业务沙盘，用于说明算法与状态流转；页面结构图是功能示意；附件D为开发版实际截图。沙盘复算与产品业务验收分别记录，真实评级、完整结果及生产签收依实际证据确认。','Caption')
doc.add_page_break()

source=(ROOT/'reports/stakeholder/business-narrative.md').read_text().splitlines();i=0
while i<len(source):
    line=source[i].strip()
    if line.startswith('## '):heading(line[3:])
    elif line.startswith('### '):heading(line[4:],2)
    elif line.startswith('!['):
        m=re.fullmatch(r'!\[([^]]+)\]\(([^)]+)\)',line);picture(m[2],m[1])
    elif line=='{{CASE_SETUP}}':
        heading('贯穿案例 远景制造 SYN-A 和 SYN-H',2)
        for text in [
            '上年收入100亿元；普通股TTM盈利10亿元；同权股本10亿股；EPS1.00元。',
            '公司分79.0、覆盖100%；A/H同权假设；1港元=0.90元人民币。',
            '第2-5节沿同一对象推演A/H、订单、财报吸收与风险；D0-D7为沙盘时点。']:
            p=para(text);p.paragraph_format.space_after=Pt(3);p.paragraph_format.keep_with_next=True
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'F0F5F7');p._p.get_or_add_pPr().append(shade)
    elif line=='{{CASE_TIMELINE}}':
        rows=[]
        for k,s in enumerate(case['sessions']):
            price=s['h_price_hkd'];pv='缺报价' if k==2 else ('无需新报价' if k==7 else fmt(D(price))+'港元\n'+fmt(pe(price))+' / '+fmt(value(price)))
            status=s['display']+(' '+s['counter'] if k<5 else '')+'\n已确认：'+s['last_confirmed']
            rows.append([s['id'],pv,fmt(q_at(k)),status,s['event']])
        table(['日','H价 / PE / 估值分','公司分','显示 / 确认','业务原因'],rows,[1.0,3.5,1.6,4.0,6.9],fs=10.3)
        para('A股演示价格持续20元，PE20、估值分33.33，普通条件下一直在区间外。D1-D5事件年龄0至4自然日，按30日半衰期衰减；D6假定有效新基准吸收影响，成长基准64、事件0，公司分79.40。','Caption')
    elif line.startswith('|'):
        headers=[x.strip() for x in line.strip('|').split('|')];rows=[];i+=2
        while i<len(source) and source[i].strip().startswith('|'):
            rows.append([x.strip() for x in source[i].strip().strip('|').split('|')]);i+=1
        i-=1
        if headers[0]=='经营维度':widths=[3.8,1.4,1.7,10.1]
        elif headers[0]=='同一经营基础':widths=[6.7,4.6,5.7]
        elif headers[0]=='变化':widths=[3.2,6.4,7.4]
        elif headers[0]=='工作角色':widths=[3.1,7.4,6.5]
        else:widths=[3.2,7.1,6.7]
        table(headers,rows,widths)
    elif line and not line.startswith('# '):para(line)
    i+=1

heading('附件A 生产访问、资源与运行费用')
para('上线需要一套能够长期运行和恢复的研究环境：域名是入口，可信登录决定谁能读和修改，持久存储保留原件与历史，监控和备份支撑日常维护。以下是首期受控运行的询价口径，容量和恢复目标需在联调中量测。')
picture('B-06.png','B-06 生产访问与恢复职责，目标方案')
table(['资源','首期询价规格','配置与维护要求'],[
 ('应用与后台','4 vCPU / 8GB；或合并主机8 vCPU / 16GB','页面、API和独立任务进程；监督重启，发布与回退。'),
 ('数据库','2 vCPU / 4-8GB，SSD约100GB','同地域私网，运行与升级权限分开；与自建/托管方案匹配。'),
 ('原件与备份','原件100-200GB；异机备份200-500GB起','应用重建保留资料；数据库一致性备份、连续日志归档及隔离恢复。'),
 ('域名与可信身份','1个企业子域；企业IdP优先','HTTPS443、证书续期、服务端身份与角色验证。'),
 ('网络与运维','带宽按资料量和并发询价','数据库和管理端不公网开放；监控任务积压、磁盘、资料新鲜度及备份。')],[3.1,6.1,7.8])
heading('预算：先看构成，再确定完整报价',2)
para('已有公开价目参考为主机96美元/月、每日主机备份28.80美元/月、对象存储5美元/月，三项合计129.80美元/月。按预算假设1美元=7.2元，约934.56元/月；连续12个月为1,557.60美元。这是三项成本参考，不是整套系统报价。')
para('预算按已明确基础设施计基准价；身份、数据/模型、增量流量、恢复/高可用与维护人工待用途和规模明确后询价，建设工时另计。')
para('容灾目标为RPO≤15分钟、RTO≤4小时。主机镜像作为补充，数据库采用一致性基线备份、WAL连续归档及异机持久保存，按目标验证隔离恢复。是否配置热备依可用性和成本选择；日志归档延迟、完整恢复链及演练结果共同决定目标是否达成。')
para('同域名页面与/api入口便于使用。Java替换后，域名、数据库、存储、备份和身份的资源类别基本相同；迁移、适配和结果一致性验收会增加工时，不能仅靠换语言承诺降低租赁费。')

heading('附件B 技术资产与框架采用依据')
para('核心业务采用通用框架和可自托管平台，研究规则、数据库和历史记录具备自主管理基础，可按目标建设私有部署，降低单一平台绑定。自主权仍受软件许可、数据用途及外部服务合同约束；当前采用与候选分别标识。软件授权0不包含运行、维护及商业支持，扩展按需求逐步选择。')
fw=[
 ('FW-01 React','已采用：组织七个业务页面，复用组件与状态，统一研究和管理体验。','MIT；软件授权0，页面开发维护及托管另计。'),
 ('FW-02 Python / FastAPI','已采用：统一业务接口和资料处理，复用现有数据生态；策略由后端执行。','Python开源许可/FastAPI MIT；授权0，计算和研发维护另计。'),
 ('FW-03 SQLAlchemy / Alembic','已采用：数据库事务及结构升级有版本，结果、审计和任务一起保存。','MIT；授权0，迁移开发、备份及维护另计。'),
 ('FW-04 PostgreSQL','已采用：保存可信事实、不可变历史和研究，提供精确数值与事务约束。','PostgreSQL License；授权0，自建/托管、磁盘和恢复另计。'),
 ('FW-05 Linux / 容器','生产准备：固定发布制品，监督常驻进程，便于恢复；部署形式待定。','标准开源版通常无单独授权费，主机和运维支持另计。'),
 ('FW-06 Nginx / Caddy','生产候选二选一：统一HTTPS、路由、证书与请求限制。','开源版BSD/Apache-2.0；授权0，域名、带宽、商业版另计。'),
 ('FW-07 企业身份 / Keycloak','待接入：优先复用企业身份，服务端映射角色与工作区。','企业服务按合同；Keycloak Apache-2.0授权0，主机和维护另计。'),
 ('FW-08 S3兼容存储','待适配：持久保存原件及异机备份，管理保留期。','按容量、请求及流量收费；参考Spaces 5美元/月起。'),
 ('FW-09 监控与恢复平台','待选：发现来源/任务问题，验证备份可恢复，支持故障处置。','云服务或自建方案待报价，含计算、存储、维护与演练。'),
 ('FW-10 Redis / Celery','扩展候选：复杂队列、调度与缓存；当前由数据库和独立进程承接任务。','Celery BSD授权0；Redis许可随版本变化，托管及维护另计。'),
 ('FW-11 Elasticsearch','后续全文方案已确定、未接入：公告/资料匹配、高亮及筛选；采用全文生态并保留向量扩展。','默认发行物ELv2；Basic基础可免费，进阶功能/托管/支持及资源另计。'),
 ('FW-12 pgvector','语义可选替代：先评估Elasticsearch向量能力；确有增益才采用，不默认双套索引。','PostgreSQL License授权0；文本向量生成、存储、索引及维护另计。'),
 ('FW-13 外部模型平台','已确认目标：对获准资料提出受限建议，以证据和政策决定生效。','真实API待接通；服务/模型未定，调用与接入维护费用待报价。'),
 ('FW-14 Java / Spring Boot','后端备选：适合现有Java团队；基础资源复用，业务口径保持一致。','Temurin JDK/Spring社区版可无单独授权费，迁移、支持和验收另计。')]
table(['框架/平台','作用与采用依据','许可及成本'],fw,[3.8,7.5,5.7],fs=10.5)
heading('全文与语义的选择依据',2)
para('全文搜索解决公告与研究资料的关键词查找，后续统一采用Elasticsearch。选择其全文检索生态并保留同平台向量扩展路径；中文分词、权限隔离、召回及容量仍需实测。公司代码、日期和财务以精确查询为准。')
para('语义召回先评估Elasticsearch向量能力；有增益才启用，pgvector为复用数据库的替代候选。搜索候选回查原文，不直接计分或入选。')
para('默认发行物ELv2，部分免费源码另有AGPLv3等选项。Basic基础可免费，进阶能力、托管及支持另询价；搜索节点、索引磁盘、向量生成和维护未含在129.80美元示例内。')
para('业务测试矩阵覆盖数值口径、状态确认、权限及审计链，自研规则按案例验证；通用框架提供工程基础。完整投研闭环以业务结果回读、失败路径和用户任务验收为交付条件。','Caption')

heading('附件C 数据服务、费用与资料依据')
table(['数据/服务','采用依据与阶段','费用与用途条件'],[
 ('Wikimedia资料','已有有限公司文字资料，可回查原文。','公开API，无数据订阅采购；署名、存储及用途条件另核。'),
 ('BaoStock','已有A股参考日线与财务观察；不能代替完整法定财报。','公开免费类别；正式多人托管/商业用途及保障待核。'),
 ('发行人法定资料','提供原始科目、股本及币种证据，优先核对法定原文。','原件整理有人工成本；本机个人用途不自动覆盖企业再分发。'),
 ('ECB / HKMA参考汇率','当前ECB双币种22日参考汇率已接入，匹配日期及实际知悉时间。','ECB免费API，存储/适配另计；参考汇率不是香港收盘同时点价。'),
 ('EODHD','已有身份/行情适配；本项目实际市场核验未取得港交所挂牌。','个人公开参考EOD19.99、Fundamentals59.99、All-in-One99.99美元/月；不是企业多人报价。'),
 ('富途参考行情 / 正式源','两股各22条日线已按本人确认用途接入；正式最终性仍待独立政策。','未以既有个人账号推定生产许可；账户/API权限、多人留存展示与正式数据报价另核。')],[3.4,6.8,6.8],fs=10.5)
para('资料与模型按覆盖对象、使用权和调用量独立询价，数据留存、展示及外部分析逐项核定。预算保留待报价项，确保每笔支出能对应明确能力、含量与合同条件。','Caption')
p=para('公开依据（基础价目2026-10-06；Elastic资料及报告2026-10-07）：','Caption')
for label,url in [('主机价目','https://www.digitalocean.com/pricing/droplets'),('备份价格','https://docs.digitalocean.com/products/backups/details/pricing/'),('对象存储','https://docs.digitalocean.com/products/spaces/details/pricing/'),('数据套餐','https://eodhd.com/pricing'),('Elastic许可','https://www.elastic.co/pricing/faq/licensing'),('Elastic订阅','https://www.elastic.co/subscriptions'),('向量检索','https://www.elastic.co/docs/solutions/search/vector/dense-vector')]:
    link(p,label,url);p.add_run('  ')
p=para('依据：02/11/12/14及现有配置；阶段与页面按PROJECT、验证记录及前端核对。完整来源版本、价格条件和逐文件哈希留存制作记录。','Caption');p.paragraph_format.keep_with_next=False


# Actual app screenshots preserve the captured pixels; paired narrow views stay readable on A4.
shots=json.loads((ROOT/'reports/stakeholder/screenshots-v0.6.json').read_text())['images']
doc.add_page_break()
heading('附件D 实际页面与日常工作')
para('以下为2026-10-07本机开发版窄屏实拍，使用已有研究身份，只读查看。实际资料、参考结果及待补状态按画面保留；与远景制造合成沙盘分别解读。系统支持页待按运营角色补图，不用示意图冒充实拍。','Caption')
def shotpair(left,right,title,caption,description):
    heading(title,2)
    pair=doc.add_table(rows=1,cols=2);pair.autofit=False;pair.alignment=WD_TABLE_ALIGNMENT.CENTER
    pair.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
    for c,entry in zip(pair.rows[0].cells,[left,right]):
        c.width=Cm(8.5);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
        p=c.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
        p.paragraph_format.line_spacing=1.0;p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0)
        path=ROOT/entry['file'];assert hashlib.sha256(path.read_bytes()).hexdigest()==entry['sha256']
        sh=p.add_run().add_picture(str(path),width=Cm(7.5));sh._inline.docPr.set('descr',entry['id']+' '+entry['page']+' 开发版实际窄屏截图')
    p=para(caption,'Caption');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    para(description)
shotpair(shots[0],shots[1],'从今日概览进入公司研究','左 S-01 今日概览  右 S-02 公司研究 首屏局部',
    '研究员先从概览选择对象，再查看公司依据和缺口。画面中的85.1是已覆盖维度的参考分，覆盖度50%，完整经营质量仍待评估；不能与沙盘79.0完整评分直接比较。公司页继续向下可查A/H参考估值、财务原值及经营判断。')
doc.add_page_break()
shotpair(shots[2],shots[3],'审阅经营判断并回看历史研究','左 S-03 当前研判区域  右 S-04 研究快照',
    '研判区域显示建议、适用期、理由与反证，待确认建议尚未计分。快照页保留2026-10-06的已保存预览和比较入口，使团队回看当时输入；截图不触发生成、确认或保存。资料阅读、确认判断和发布规则是不同权限动作。')
doc.add_page_break()
shotpair(shots[4],shots[5],'审阅策略规则并组织个人研究','左 S-05 策略候选池  右 S-06 我的工作台',
    '策略页实拍为已发布版本5，经营分进入门槛75；正文沙盘示范规则为70，各按自己的版本解释。正式评估仍有未完成输入，规则已发布不等于证券已入选。工作台按证券维护自选和按公司记笔记；本次保留空态，未为展示补造记录。')
para('待补资料集中管理：运营角色的系统支持画面、Elasticsearch中文搜索及语义效果对比、生产域名与企业登录、正式评估完整路径、恢复演练和完整报价。按联调结果补入同一报告，不预填成功截图。','Caption')

path=OUT/f'value-investment-system-report-v{VERSION}.docx';doc.save(path)
evidence={'report_version':VERSION,'created_at':datetime.now(ZoneInfo('Asia/Shanghai')).isoformat(timespec='seconds'),'content_revision':REV,'source_hashes':HASHES,'source_changed_during_build':[p for p,h in HASHES.items() if hashlib.sha256((ROOT/p).read_bytes()).hexdigest()!=h],
 'docx':str(path.relative_to(ROOT)),'docx_sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'figures':figures,'diagram_count':len(figures),'actual_screenshots':len(shots),'screenshot_manifest':'reports/stakeholder/screenshots-v0.6.json','screenshot_metadata':shots,'oversized_screenshot_placeholders':0,'component_count':len(fw),'inventory_level':'framework_and_platform','source_count':6,'chapter_forced_page_breaks':False,'caption_keep_with_next':True,'compact_tables_kept_together':['经营维度','工作角色','日'],'business_text_source':'reports/stakeholder/business-narrative.md','case_fixture':'reports/stakeholder/case-study.json','case_arithmetic':{'eps':str(EPS),'quality_base':str(Q),'current_event':str(EVENT),'weights':scoring['dimension_weights'],'days':[{'day':s['id'],'q':fmt(q_at(i)),'pe':fmt(pe(s['h_price_hkd'])) if s['h_price_hkd'] else None,'v':fmt(value(s['h_price_hkd'])) if s['h_price_hkd'] else None} for i,s in enumerate(case['sessions'])]},'network':False,'database_access':False,'product_tests_run':False,'qa_status':'awaiting_render_and_visual_inspection'}
(ROOT/'review/stakeholder-report-build-evidence.json').write_text(json.dumps(evidence,ensure_ascii=False,indent=2)+'\n')
print(json.dumps({'docx':str(path),'figures':len(figures),'frameworks':len(fw),'source_revision':REV},ensure_ascii=False))
