"""Build the local stakeholder report. No network, application imports or database access."""
from pathlib import Path
from datetime import datetime
from zoneinfo import ZoneInfo
import hashlib, json, math, re, subprocess
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
SOURCE_REVISION=subprocess.check_output(['git','rev-parse','HEAD'],cwd=ROOT,text=True).strip()
SOURCE_FILES=['PROJECT.md','README.md','CONTEXT.md','docs/02-business-design.md',
 'docs/11-templates-automation-and-roles.md','docs/12-time-numerics-and-corrections.md',
 'docs/14-domain-and-business-flows.md','versions/v0.0.1/VERIFICATION.md',
 'review/REAL-DATA-AND-JUDGMENTS.md','review/REAL-GAP-CLOSURE.md','review/FUTU-INTEGRATION.md',
 'reports/stakeholder/business-narrative.md','reports/stakeholder/components-and-costs.md',
 'reports/stakeholder/frameworks-and-costs.md','research/report-component-cost-basis.md',
 'research/official-fx-hk-recheck.md','backend/pyproject.toml','backend/uv.lock','frontend/package.json']
SOURCE_HASHES={p:hashlib.sha256((ROOT/p).read_bytes()).hexdigest() for p in SOURCE_FILES}
OUT=ROOT/"output/documents"
ASSETS=ROOT/"reports/stakeholder/figures"
OUT.mkdir(parents=True,exist_ok=True)
ASSETS.mkdir(parents=True,exist_ok=True)
FONT_TMP=ROOT/'tmp/stakeholder-report'
FONT_TMP.mkdir(parents=True,exist_ok=True)
(FONT_TMP/'fonts.conf').write_text('''<?xml version="1.0"?>
<!DOCTYPE fontconfig SYSTEM "urn:fontconfig:fonts.dtd">
<fontconfig>
<dir>/System/Library/Fonts</dir>
<dir>/System/Library/Fonts/Supplemental</dir>
<dir>/Users/kyle/.cache/codex-runtimes/codex-primary-runtime/dependencies/native/libreoffice-headless/libreoffice/LibreOfficeDev.app/Contents/Resources/fonts/truetype</dir>
<cachedir>'''+str(FONT_TMP/'font-cache')+'''</cachedir>
</fontconfig>''')
FONT="/System/Library/Fonts/Hiragino Sans GB.ttc"
FONTNAME="Hiragino Sans GB"
BLUE="173E56"; GREY="D9D9D9"; PALE="F3F6F8"
figures=[]
def font(size):
    return ImageFont.truetype(FONT,size)
def wrap(text,f,maxw,draw):
    lines=[]
    for part in text.split("\n"):
        line=""
        for ch in part:
            if draw.textlength(line+ch,font=f)>maxw and line:
                lines.append(line);line=ch
            else:line+=ch
        lines.append(line)
    return lines
def diagram(fid,nodes,edges,size):
    im=Image.new("RGB",size,"white");d=ImageDraw.Draw(im)
    for points,dashed in edges:
        for a,b in zip(points,points[1:]):
            if dashed:
                dist=math.dist(a,b)
                for t in range(0,int(dist),22):
                    u=t/dist;v=min(t+12,dist)/dist
                    d.line((a[0]+(b[0]-a[0])*u,a[1]+(b[1]-a[1])*u,
                            a[0]+(b[0]-a[0])*v,a[1]+(b[1]-a[1])*v),fill="#7C8992",width=5)
            else:d.line([a,b],fill="#637986",width=5)
        a,b=points[-2:];ang=math.atan2(b[1]-a[1],b[0]-a[0])
        d.polygon([b,(b[0]-21*math.cos(ang-.45),b[1]-21*math.sin(ang-.45)),
                     (b[0]-21*math.cos(ang+.45),b[1]-21*math.sin(ang+.45))],fill="#637986")
    for rect,title,sub,state in nodes:
        x,y,w,h=rect
        fill={"current":"#EDF4F7","target":"#F8F3E9","pending":"#F7F0DC","neutral":"#F3F6F8"}[state]
        d.rounded_rectangle((x,y,x+w,y+h),radius=16,fill=fill,outline="#8DA2AD",width=3)
        f=font(42);s=font(35)
        lines=wrap(title,f,w-38,d);small=wrap(sub,s,w-38,d) if sub else []
        ht=len(lines)*57+len(small)*44+(12 if small else 0)
        yy=y+(h-ht)/2
        for line in lines:
            d.text((x+(w-d.textlength(line,font=f))/2,yy),line,font=f,fill="#172E3B");yy+=57
        yy+=12 if small else 0
        for line in small:
            d.text((x+(w-d.textlength(line,font=s))/2,yy),line,font=s,fill="#4B606C");yy+=44
    path=ASSETS/(fid+".png");im.save(path,dpi=(300,300))
    figures.append({"figure_id":fid,"file":str(path.relative_to(ROOT)),"pixel_size":list(size),
                    "kind":"locally_drawn_business_diagram","not_a_screenshot":True})
    return path
N=lambda x,y,w,h,t,s="",k="current":((x,y,w,h),t,s,k)
E=lambda *pts,dash=False:(list(pts),dash)

diagram("F-01",[
 N(40,90,390,210,"选研究对象","先看结论与缺口"),N(480,90,390,210,"进入公司档案","经营与证券分开"),
 N(920,90,390,210,"回查依据","核对财务与原文"),N(1360,90,390,210,"补充或复核判断","有证据且有权限"),
 N(1360,460,390,210,"保存当次研究","保留输入与缺口"),N(920,460,390,210,"比较前后结果","看变化及原因"),
 N(480,460,390,210,"继续跟踪","关注证券与策略")],
 [E((430,195),(480,195)),E((870,195),(920,195)),E((1310,195),(1360,195)),
  E((1555,300),(1555,460)),E((1360,565),(1310,565)),E((920,565),(870,565))],(1800,760))

diagram("F-02",[
 N(40,40,500,160,"材料进入档案","出处 正文 取得时间"),
 N(650,40,500,160,"确认研究对象","关联真实公司或证券"),
 N(1260,40,500,160,"归并同一事实","五篇转载只算一件事"),
 N(1260,290,500,160,"提出有证据的判断","对经营有何含义"),
 N(650,290,500,160,"核对已发布政策","证据 适用期 冲突检查"),
 N(40,560,720,180,"判断有效：进入评价","符合政策或有效人工决定"),
 N(1040,560,720,180,"判断未定：继续研究","保留材料与缺口 不计有效分")],
 [E((540,120),(650,120)),E((1150,120),(1260,120)),
  E((1510,200),(1510,290)),E((1260,370),(1150,370)),
  E((900,450),(900,505),(400,505),(400,560)),
  E((900,505),(1400,505),(1400,560))],(1800,780))

diagram("F-03",[
 N(650,45,500,175,"浏览器与React页面","七个功能区域"),
 N(650,300,500,185,"FastAPI业务接口","本地开发身份"),
 N(650,850,500,200,"PostgreSQL事实真源","资料 历史 研究 审计"),
 N(40,300,490,185,"显式资料导入与解析","公开数据接口 财报解析"),
 N(40,850,490,200,"本机原件与快照","原PDF 内容hash 定位"),
 N(1270,600,490,185,"独立后台任务","任务恢复 避免重复执行"),
 N(1270,45,490,175,"固定模板与业务合同","策略 数值 时间口径")],
 [E((900,220),(900,300)),E((900,485),(900,850)),
  E((285,485),(285,850)),E((530,392),(585,392),(585,770),(770,770),(770,850)),
  E((1270,692),(1200,692),(1200,950),(1150,950)),
  E((1270,132),(1200,132),(1200,392),(1150,392)),
  E((1515,220),(1515,600))],(1800,1140))

diagram("F-04",[
 N(650,40,500,165,"授权用户与域名入口","DNS HTTPS 443","target"),
 N(650,295,500,180,"应用API与静态UI","反向代理后运行","target"),
 N(1270,295,490,180,"可信登录与用户映射","公网访问前置条件","target"),
 N(40,575,490,180,"后台Worker","独立监督与调度","target"),
 N(650,855,500,200,"私网PostgreSQL","应用与迁移角色分离","target"),
 N(40,855,490,200,"持久原件存储","授权留存与生命周期","target"),
 N(650,1220,500,180,"异机备份与恢复","物理基线 WAL 回读","target"),
 N(1270,855,490,200,"后续接入能力","检索 队列 外部模型","target")],
 [E((900,205),(900,295)),E((1270,385),(1150,385),dash=True),
  E((900,475),(900,855)),
  E((530,665),(590,665),(590,955),(650,955)),
  E((650,385),(285,385),(285,575),dash=True),
  E((650,445),(570,445),(570,805),(285,805),(285,855)),
  E((900,1055),(900,1220)),E((285,1055),(285,1310),(650,1310)),
  E((1150,440),(1200,440),(1200,770),(1515,770),(1515,855),dash=True)],(1800,1470))

diagram("F-05",[
 N(40,40,520,160,"研究预览","随时看结果与缺口","target"),
 N(640,40,520,160,"核验正式评估时点","合法日历 最终收盘依据","target"),
 N(1240,40,520,160,"固定当次依据","资料 判断 规则和缺口","target"),
 N(40,300,520,215,"普通条件变化","已有基线后\n两个相邻有效收盘确认","target"),
 N(640,300,520,215,"关键资料缺失","待评估 保留已确认状态\n打断连续确认","target"),
 N(1240,300,520,215,"已确认重大硬风险","按适用范围优先处理\n不等普通确认次数","target"),
 N(640,620,520,160,"留下结果与变化解释","回查当时依据","target")],
 [E((560,120),(640,120)),E((1160,120),(1240,120)),
  E((1500,200),(1500,250),(300,250),(300,300)),
  E((900,250),(900,300)),E((1500,250),(1500,300)),
  E((300,515),(300,565),(780,565),(780,620)),
  E((900,515),(900,620)),E((1500,515),(1500,565),(1020,565),(1020,620))],(1800,820))

diagram("F-06",[
 N(40,60,390,225,"资料与证据","保留发生了什么\n以及原始出处"),
 N(480,60,390,225,"研究与判断","确认涉及谁\n意味着什么"),
 N(920,60,390,225,"质量与估值","评价企业经营\n分别评价证券"),
 N(1360,60,390,225,"策略与跟踪","比较已发布条件\n解释状态变化"),
 N(40,465,1710,170,"历史与复核","保存每次依据与结果，比较变化，处理修订和原错误")],
 [E((430,172),(480,172)),E((870,172),(920,172)),E((1310,172),(1360,172)),
  E((235,285),(235,465)),E((675,285),(675,465)),
  E((1115,285),(1115,465)),E((1555,285),(1555,465))],(1800,700))

diagram("F-07",[
 N(650,40,500,160,"同一家公司","企业经营研究对象"),
 N(400,280,1000,185,"共享经营档案","财报 经营事实 有效判断 经营质量"),
 N(40,635,760,235,"A股证券","A股价格 币种 适用股数\nA股估值 → A股策略结果"),
 N(1000,635,760,235,"H股证券","H股价格 币种 适用股数\nH股估值 → H股策略结果")],
 [E((900,200),(900,280)),E((900,465),(900,535),(420,535),(420,635)),
  E((900,535),(1380,535),(1380,635))],(1800,920))

diagram("F-08",[
 N(40,40,730,160,"公司经营依据","财务 + 带证据的经营判断"),
 N(1030,40,730,160,"每只证券的估值依据","价格 币种 股数 财务 汇率"),
 N(40,290,730,170,"经营质量及覆盖度","经营好不好 有多少有效依据"),
 N(1030,290,730,170,"证券估值及缺口","该价格是否符合估值条件"),
 N(450,580,900,170,"策略综合判断","质量 覆盖 估值 财务条件 风险")],
 [E((405,200),(405,290)),E((1395,200),(1395,290)),
  E((405,460),(405,520),(650,520),(650,580)),
  E((1395,460),(1395,520),(1150,520),(1150,580))],(1800,790))

def placeholder(sid,title):
    im=Image.new("RGB",(1800,420),"#F6F7F8");d=ImageDraw.Draw(im)
    for x in range(24,1776,36):
        d.line((x,24,min(x+20,1776),24),fill="#A9B4BB",width=3)
        d.line((x,396,min(x+20,1776),396),fill="#A9B4BB",width=3)
    for y in range(24,396,36):
        d.line((24,y,24,min(y+20,396)),fill="#A9B4BB",width=3)
        d.line((1776,y,1776,min(y+20,396)),fill="#A9B4BB",width=3)
    for text,yy,size,color in [(sid+" 实际系统截图待补",100,50,"#4D626E"),
                              (title,188,39,"#4D626E"),
                              ("占位区域  未采集画面  版本 角色 日期待登记",278,33,"#71828B")]:
        f=font(size);d.text(((1800-d.textlength(text,font=f))/2,yy),text,font=f,fill=color)
    path=ASSETS/(sid+"-placeholder.png");im.save(path,dpi=(300,300));return path

def clean(s):
    s=re.sub(r"\[([^\]]+)\]\([^)]+\)",r"\1",s)
    return s.replace("**","").replace(chr(96),"").replace("…","至")
doc=Document();sec=doc.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=Cm(1.85);sec.bottom_margin=Cm(1.85)
sec.left_margin=Cm(2);sec.right_margin=Cm(2)
sec.header_distance=Cm(.7);sec.footer_distance=Cm(.8)
for name,size in [("Normal",11.5),("Title",28),("Subtitle",14),("Heading 1",19),("Heading 2",13.5),("Caption",9.5)]:
    st=doc.styles[name];st.font.name="Arial";st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
    st._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"),FONTNAME)
    st.paragraph_format.line_spacing=1.22
    st.paragraph_format.space_after=Pt(7)
for element in doc.styles._element.xpath('.//w:rFonts'):
    for attr in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme','csTheme']:
        element.attrib.pop(qn('w:'+attr),None)
    element.set(qn('w:eastAsia'),FONTNAME)
for element in doc.styles._element.xpath('.//w:pBdr'):
    element.getparent().remove(element)
doc.styles["Normal"].paragraph_format.widow_control=True
doc.styles['Normal'].font.bold=False
doc.styles['Caption'].font.bold=False
doc.styles['Caption'].font.italic=False
doc.styles["Caption"].paragraph_format.space_after=Pt(7)
doc.core_properties.title="价值投资研究系统 业务方案 系统说明与开发进展"
doc.core_properties.subject="需求方审阅报告 含截图及未完成内容占位"
doc.core_properties.author="项目团队"
doc.core_properties.keywords="价值投资 系统说明 开发进展 资源 组件 费用"
head=sec.header.paragraphs[0];head.text="价值投资研究系统  |  需求方审阅"
for run in head.runs:run.font.size=Pt(9);run.font.color.rgb=RGBColor(0,0,0)
foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
foot.add_run("版本 0.3  |  2026-10-06  |  系统开发中  |  ")
field=OxmlElement("w:fldSimple");field.set(qn("w:instr"),"PAGE");foot._p.append(field)
for run in foot.runs:run.font.size=Pt(9)
sections=[]
def start(title,intro=None):
    sections.append(title)
    heading=doc.add_heading(title,level=1)
    heading.paragraph_format.page_break_before=True
    if intro:para(intro)
def para(text,style=None):
    return doc.add_paragraph(clean(text),style=style)
def sub(title):doc.add_heading(title,level=2)
def bullets(items):
    for s in items:
        p=para("• "+s);p.paragraph_format.space_after=Pt(6)
def table(headers,rows,widths=None,fs=10.5):
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    widths=widths or [17/len(headers)]*len(headers)
    for cell,w in zip(t.rows[0].cells,widths):cell.width=Cm(w)
    for col,w in zip(t.columns,widths):col.width=Cm(w)
    pr=t._tbl.tblPr;b=OxmlElement("w:tblBorders")
    for side in ["top","left","bottom","right","insideH","insideV"]:
        a=OxmlElement("w:"+side);a.set(qn("w:val"),"single");a.set(qn("w:sz"),"4");a.set(qn("w:color"),GREY);b.append(a)
    pr.append(b)
    for j,value in enumerate(headers):t.rows[0].cells[j].text=value
    hpr=t.rows[0]._tr.get_or_add_trPr();repeat=OxmlElement("w:tblHeader");hpr.append(repeat)
    for values in rows:
        cells=t.add_row().cells
        for j,value in enumerate(values):cells[j].text=clean(str(value));cells[j].width=Cm(widths[j])
    for i,row in enumerate(t.rows):
        rp=row._tr.get_or_add_trPr();cant=OxmlElement("w:cantSplit");rp.append(cant)
        for j,cell in enumerate(row.cells):
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp=cell._tc.get_or_add_tcPr();shade=OxmlElement("w:shd")
            shade.set(qn("w:fill"),BLUE if i==0 else ("FFFFFF" if i%2 else PALE));cp.append(shade)
            margins=OxmlElement("w:tcMar")
            for side in ["top","bottom","left","right"]:
                n=OxmlElement("w:"+side);n.set(qn("w:w"),"95" if side in ["top","bottom"] else "115");n.set(qn("w:type"),"dxa");margins.append(n)
            cp.append(margins)
            for p in cell.paragraphs:
                p.paragraph_format.line_spacing=1.18;p.paragraph_format.space_after=Pt(1);p.paragraph_format.space_before=Pt(1)
                if j==0 and len(headers)>2 and widths[0]<3:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    r.font.size=Pt(fs);r.font.name="Arial";r._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"),FONTNAME)
                    r.font.bold=i==0;r.font.color.rgb=RGBColor.from_string("FFFFFF" if i==0 else "000000")
    para("")
    return t
def picture(path,caption,width=17):
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after=Pt(3);p.paragraph_format.keep_with_next=True
    run=p.add_run();shape=run.add_picture(str(path),width=Cm(width))
    shape._inline.docPr.set("descr",caption)
    c=para(caption,"Caption");c.alignment=WD_ALIGN_PARAGRAPH.CENTER

sections.append("封面")
doc.add_paragraph("价值投资研究系统",style="Title")
doc.add_paragraph("业务架构与研究流程说明",style="Subtitle")
para("需求方审阅报告",style="Heading 2")
para("报告版本 0.3   编制日期 2026-10-06")
para("系统开发中，完整必选验收尚未通过。")
sub("从资料到结论，从结论回到依据")
para("本系统把公司研究、经营判断、证券估值和策略跟踪连接起来，让用户知道当前结论基于什么、哪些部分还不确定，以及相较上次为什么发生变化。")
para("报告先解释业务职责与对象关系，再讲一份材料如何形成研究结果、结果如何持续跟踪，并用“示例制造”串起完整过程。技术框架、生产资源和费用放在后部，供实施和采购复核。")
sub("怎样阅读")
para("前部九章回答系统做什么、如何运作、人与系统如何分工。随后六组界面说明把业务路径对应到主要页面，实际截图均保留占位。最后说明当前技术实现、生产准备、框架采用依据与费用。")
sub("设计与当前实现分别说明")
para("业务架构和流程描述已确认设计；当前仍只完成部分真实研究能力，不能据此认定完整流程或生产服务已交付。第09章集中说明阶段事实和缺口。示例是合成情景，不是真实公司结果。")
para("实际截图、完整报价、正式业务结果和未完成验收均明确标注待补。报告仅本地交付，不产生交易指令。")
para("内容快照日期为2026-10-06；项目修订及引用材料在附件和制作记录登记。", "Caption")

# Narrative is the single source of truth for the business explanation.
for line in (ROOT/"reports/stakeholder/business-narrative.md").read_text().splitlines():
    if line.startswith("## "):start(line[3:])
    elif line.startswith("### "):sub(line[4:])
    elif line.startswith("!["):
        match=re.fullmatch(r"!\[([^]]+)\]\(([^)]+)\)",line)
        if not match:raise ValueError("Invalid figure directive: "+line)
        picture(ASSETS/Path(match[2]).name,match[1])
    elif line and sections[-1]!="封面" and not line.startswith("# "):para(line)

screens=[
 ("界面说明 概览与公司研究",[
  ("S-01","今日概览","看当前对象、最近研究与关键缺口，决定先研究哪家公司。","公司入口、待补事项和最近研究是阅读重点；对象数量不代表完整分析覆盖。"),
  ("S-02","经营质量与待补依据","看经营维度的证据、覆盖和待评估项，再回查材料或补充判断。","缺依据时显示待评估；待审建议与已生效判断分别显示，分数和覆盖度并列。")]),
 ("界面说明 证券与财务依据",[
  ("S-03","A/H证券卡片","同一公司不同证券独立展示价格、币种与估值输入。","参考价与正式价区分；港股未取到价格时保留缺失提示。"),
  ("S-04","财务与估值依据","追溯科目期间、原文定位、标准计算口径及股份有效日期。","财务指标和估值输入分别核对；关键分量或当前股数依据缺失时，保留具体缺口。")]),
 ("界面说明 原文与研判",[
  ("S-05","固定原文阅读","核对当次修订正文、出处和定位，再返回此前页面。","结论引用当次固定原文；读者按授权查看正文，资料后来变化不替换旧研究依据。"),
  ("S-06","资讯与研判","把原文事实转为有适用时期的判断，审阅建议理由与局限。","当前7项建议仍待审、未校准，不进入有效评分；外部模型真实调用仍待接通。")]),
 ("界面说明 快照与策略",[
  ("S-07","研究快照与历史比较","读取当次固定输入、结果阶段和前后变化原因。","当前真实研究仍不完整；保留当次依据及缺口，历史比较说明变化来自哪个环节。"),
  ("S-08","策略候选池","查看规则版本、预览与正式结果，以及状态变化解释。","当前真实正式封存结果为0；研究预览与正式入选、退出状态分别说明。")]),
 ("界面说明 规则与个人工作台",[
  ("S-09","公司评分规则","看基础、行业和企业配置来源，理解预览与发布职责。","企业差异在基础和行业规则上补充；当前生物医药专用模板尚未发布。"),
  ("S-10","我的工作台","维护个人关注的A/H证券和研究上下文，持续跟进。","A股与H股可以分别关注；个人笔记按可见权读取，具体页面截图待补。")]),
 ("界面说明 系统支持与缺失状态",[
  ("S-11","系统支持","由获授权人员查看来源、任务积压和故障说明。","数据运营处理来源与任务问题，管理员维护账号和运行连接；研究与发布仍按职责授权。"),
  ("S-12","缺数据或失败状态","理解无法完成的原因和可继续采取的下一步。","页面说明缺什么、影响哪个结果及怎样继续；缺数据不解释成零分或自动退出。")])]
for title,items in screens:
    start(title,"以下为实际截图占位，当前没有采集系统画面。")
    for sid,name,purpose,note in items:
        sub(sid+" "+name)
        picture(placeholder(sid,name),sid+" 素材状态 待采集",width=17)
        para(purpose);para(note,"Caption")

start("当前实现架构","采用模块化单体，先用少量运行进程贯通业务职责。")
picture(ASSETS/"F-03.png","F-03 当前实现架构  本机研究版本  未部署生产",width=16.4)
table(["职责","当前组件与采用依据"],[
 ("页面与接口","React提供业务页面；Python/FastAPI承接统一接口与业务逻辑，前端不另算策略。"),
 ("事实与事务","PostgreSQL保存修订、研究、审计与任务；SQLAlchemy组织同一事务。"),
 ("资料与任务","Python处理公开数据与财报；独立后台进程承接当前任务，Redis/Celery仍是扩展候选。")],[3.2,13.8])
para("当前使用本地开发身份，不能直接作为公网登录。事实数据库是业务真源，检索与缓存接入后仍是可重建投影。","Caption")

start("生产访问目标架构","同一域名提供页面和API，入口、身份、持久化与恢复共同构成上线基础。")
picture(ASSETS/"F-04.png","F-04 生产目标架构  尚未实施  虚线表示身份或扩展对接",width=15.6)
para("可信登录、用户映射和服务器端权限是公网访问前置条件。S3、Redis/Celery、OpenSearch、pgvector和外部模型按相应能力接入；不把候选组件画成已上线服务。","Caption")
para("目标恢复能力需要物理备份、WAL归档与隔离回读。仅订购每日主机镜像不能证明15分钟RPO或4小时RTO。","Caption")

start("生产资源与域名配置","以下规格用于首期受控试运行询价，未压测，不构成容量承诺。")
table(["资源","起步建议","配置要求"],[
 ("应用计算","4 vCPU / 8 GB","API、静态UI与独立Worker；Linux受支持版本，监督与自动重启。"),
 ("PostgreSQL","2 vCPU / 4-8 GB；SSD约100 GB","同地域私网，自建或托管择一；运行与迁移权限分离。"),
 ("原件持久卷","100-200 GB，可扩容","保存原PDF、快照、hash和定位；重建应用不丢原件。"),
 ("异机备份","200-500 GB起步","不同故障域；物理全备/WAL和必要原件，覆盖量随保留策略核算。"),
 ("域名与DNS","1个企业子域","例如research.example.com；控制权、续费与解析负责人明确。"),
 ("HTTPS入口","1个反向代理或托管入口","Nginx或Caddy等择一；证书续期、到期监测和443访问。"),
 ("可信身份","复用企业IdP优先","服务端验证身份，映射五角色和工作区；禁止自报开发身份。"),
 ("网络与运维","稳定入口；带宽按负载询价","PG和应用管理端口不公网开放；SSH限来源或私网。"),
 ("监控与支持","云监控或受控自建","CPU、磁盘、任务积压、证书、来源新鲜度与备份状态。")],[3.0,5.0,9.0],fs=10.3)
sub("域名访问方式")
para("推荐 https://research.example.com/ 提供页面，/api/ 提供接口。同源减少跨域配置；选择子目录部署需先改造现有绝对资源路径。")
para("单机合并可用8 vCPU / 16 GB起步估算；检索和全量历史接入后需重测。部署地域、适用备案、来源远程留存/多人使用权和恢复方案仍待确定。","Caption")

start("费用构成与预算占位","软件授权费、基础设施、数据模型与人工分别核算。")
table(["费用项","当前口径"],[
 ("软件许可","主要框架的标准开源发行版在相应条件内授权费为0；商业版本与支持另核实。"),
 ("基础设施","计算、数据库、存储、备份、网络、身份和监控；以下只列三项公开参考。"),
 ("数据与模型","港股覆盖、商业/多人许可、模型与embedding费用待正式方案，不能写0。"),
 ("建设与维护","研发、测试、资料规范化、部署恢复和运维工时；完整工作量与单价待核。")],[3.8,13.2])
sub("官方公开价目示例")
table(["项目","USD每月"],[
 ("DigitalOcean Basic Regular共享CPU 8vCPU / 16GiB / 320GiB SSD","96.00"),
 ("每日主机备份 主机费的30%","28.80"),
 ("Spaces标准 250GiB存储和1024GiB出站共用额度","5.00"),
 ("三项参考小计","129.80")],[13.0,4.0])
para("同样月租连续12个月为1,557.60美元，不是包年折扣。仅按预算假设1美元=7.2元换算，约934.56元/月、11,214.72元/年；不是实时汇率。","Caption")
para("参考不含域名、超量磁盘/流量、数据库一致性与WAL/PITR、HA、额外身份监控、数据、模型、人工、税费等。主机备份不含附加Volumes，不能替代活跃数据库的应用级备份。","Caption")
sub("完整预算待补")
para("【报价占位】云厂商和地域、数据用途/套餐、模型调用量、恢复资源及人工工时未确定，系统完整月费与建设总价暂不填金额。")

start("组件成熟度与Java备选","采用依据是业务需要、兼容性和维护成本，成熟框架仍需本项目验证。")
table(["组件类别","判断依据与限制"],[
 ("通用技术","React、Python、FastAPI、SQLAlchemy、PostgreSQL具备生产应用基础；固定版本、权限、性能与恢复需验收。"),
 ("数据服务","公开和商业数据的字段覆盖、可用性、许可和服务保障单独核查；框架免费不授予数据商用权。"),
 ("自研业务功能","评分、策略、封存和避免重复执行的机制按本项目合同验收，不能靠框架流行程度证明。"),
 ("目标扩展","Redis/Celery、搜索、身份、模型等区分当前与候选；未接入不计为已有能力。")],[3.7,13.3])
sub("换成Java的资源类别基本相同")
table(["复用或替换","作用与费用影响"],[
 ("域名、PG、存储、备份与身份","可以沿用；资源量、连接池和恢复仍需验证。"),
 ("Temurin JDK与Spring Boot/Security","承接API与认证；社区软件授权费可为0，商业支持和集成工时另计。"),
 ("Java业务与数据库迁移","保留数值、权限与历史口径；迁移前后业务结果须一致，数据库结构由统一方案管理。"),
 ("资料适配与Python保留","资料处理生态无合适替代时可保留Python采集服务，仍有运行和维护费用。")],[5.2,11.8],fs=10.4)
para("尚未决定或实施Java迁移。新增成本主要来自研发、适配和验收，不能承诺换语言即降低租赁费；Python和Java备选不叠加到当前小计。","Caption")

start("下一阶段与缺失材料占位","先补足真实业务依赖，再完成正式闭环、用户签收和生产验证。")
table(["阶段","需要形成的结果","完成依据"],[
 ("真实研究依据","港股价格、当前股份、有效经营判断及同日参考汇率匹配。","固定原文与口径可回查；参考汇率不冒充香港收盘时点汇率。"),
 ("正式业务闭环","批准日历、最终价格政策、封存、风险和纠错。","依原R/W/T和业务回读验收；不以build/health替代。"),
 ("用户与生产验证","用户任务签收、可信登录、发布配置和隔离恢复。","真实权限负例、业务路径与RPO/RTO量测。"),
 ("报告更新","补截图、刷新阶段状态、取得适用报价。","素材版本与图注一致，价格可复算，逐页检查。")],[3.1,7.6,6.3],fs=10.4)
sub("缺失材料按以下方式补位")
bullets(["【截图待补 S-01至S-12】实际页面、版本、角色、日期与脱敏范围登记后替换。","【业务结果待补】经营有效判断、完整估值和正式策略结果；当前保留缺口。","【部署与报价待补】地域、云厂商、账号体系、数据用途、模型政策与维护责任。","【验收待补】完整必选、用户签收和恢复演练；当前不写通过结论。"])
para("需求方可据此核对首批对象、使用成员、模板/经营判断维护责任和部署使用范围。已确认的多人、三层模板、自动政策与人工覆盖、通知后置不重新列为未决。")
para("本报告先提供可审阅版本；替换截图和补齐结果不会覆写历史研究依据。","Caption")

# Stakeholder inventory stays at framework/platform level; dependency details remain internal.
raw=(ROOT/"reports/stakeholder/components-and-costs.md").read_text()
framework_raw=(ROOT/"reports/stakeholder/frameworks-and-costs.md").read_text()
component_rows={}
data_rows={}
for line in framework_raw.splitlines():
    if re.match(r"^\| FW-\d{2} ",line):
        cells=[s.strip() for s in line.strip("|").split("|")]
        cid=re.search(r"FW-\d{2}",cells[0]).group()
        component_rows[cid]=[cells[0],"作用："+cells[1]+"；依据："+cells[2],cells[-1]]
for line in raw.splitlines():
    if re.match(r"^\| D-\d{2} ",line):
        cells=[s.strip() for s in line.strip("|").split("|")]
        data_rows[re.search(r"D-\d{2}",cells[0]).group()]=[cells[0],cells[1],cells[2]+"；条件："+cells[3]]
groups=[
 ("附件 当前主要框架",[f"FW-{i:02}" for i in range(1,5)],"主要框架已用于本机研究版本，尚未部署生产。"),
 ("附件 生产平台与基础设施",[f"FW-{i:02}" for i in range(5,10)],"生产准备或候选平台，域名与可信登录须在发布前落实。"),
 ("附件 扩展框架与Java备选",[f"FW-{i:02}" for i in range(10,15)],"待接入或备选，按实际业务需要采用，不与当前费用重复相加。")]
for title,ids,intro in groups:
    start(title,intro)
    table(["框架与平台","作用与采用依据","许可及费用"],[component_rows[i] for i in ids],[3.4,7.1,6.5],fs=10.25)
    para("授权费0以遵守对应发行版许可为条件；保留版权、许可与适用NOTICE，商业支持、基础设施和运维另计。逐项公开许可依据见来源附件。","Caption")

start("附件 数据服务与费用","来源必须同时满足实际覆盖、合法取得、留存、分析和展示用途。")
table(["来源","作用与采用依据","费用与使用条件"],list(data_rows.values()),[3.4,6.1,7.5],fs=10.1)

start("附件 版本依据与公开来源","本报告使用项目记录与公开软件许可、价目；系统数据和截图未作为原件附入。")
sub("项目依据")
para("内容修订："+SOURCE_REVISION+"；采用制作时读取的工作树，来源摘要哈希另留本机制作记录。", "Caption")
bullets(["PROJECT.md与v0.0.1/VERIFICATION：当前阶段与未完成必选。","review/REAL-GAP-CLOSURE：182项原值、ECB参考FX及剩余依赖；FUTU-INTEGRATION：223项隔离合成检查及港股真实接入前置。","review/REAL-FINANCIAL-CONTINUATION及COMPANY-09969-ADDITION：标准财务、市场和公司补充证据；7项建议仍待审。","docs/02、11、12与14：业务设计、对象关系、流程、角色、模板、自动政策、人工覆盖、时点及纠错合同。","docs/19与reports/stakeholder/frameworks-and-costs：资源建议、框架与费用口径。"])
sub("公开价目  核查于2026-10-06")
for title,url in [
 ("DigitalOcean主机","https://www.digitalocean.com/pricing/droplets"),
 ("主机备份价格","https://docs.digitalocean.com/products/backups/details/pricing/"),
 ("备份限制","https://docs.digitalocean.com/products/backups/details/limits/"),
 ("Spaces价格","https://docs.digitalocean.com/products/spaces/details/pricing/"),
 ("EODHD公开价格","https://eodhd.com/pricing")]:
    para(title+"  "+url,"Caption")
para("美元价目不含税费、支付与汇率差异；正式采购时刷新地区可售、额度、使用范围及续费条件。恢复目标与容量仍需实测。")

start("附件 框架许可与恢复依据","主要框架的软件费按对应发行版条件核实，托管与商业支持另计。")
for title,url in [
 ("React","https://github.com/facebook/react/blob/main/LICENSE"),
 ("Python","https://docs.python.org/3/license.html"),
 ("FastAPI","https://github.com/fastapi/fastapi/blob/master/LICENSE"),
 ("PostgreSQL","https://www.postgresql.org/about/licence/"),
 ("Redis版本许可","https://redis.io/legal/licenses/"),
 ("SQLAlchemy","https://github.com/sqlalchemy/sqlalchemy/blob/main/LICENSE"),
 ("Keycloak","https://github.com/keycloak/keycloak/blob/main/LICENSE.txt"),
 ("Temurin JDK","https://adoptium.net/docs/faq"),
 ("Spring Boot","https://github.com/spring-projects/spring-boot/blob/main/LICENSE.txt"),
 ("PostgreSQL连续归档与恢复","https://www.postgresql.org/docs/current/continuous-archiving.html")]:
    para(title,"Heading 2");para(url,"Caption")
para("授权费0以遵守具体发行版许可为条件。框架成熟度与本系统生产验收分别核实；固定版本、使用范围和商业服务条件在上线前锁定。","Caption")

path=OUT/"value-investment-system-report-v0.3.docx"
doc.save(path)
manifest={
 "report_version":"0.3","created_at":datetime.now(ZoneInfo("Asia/Shanghai")).isoformat(timespec="seconds"),
 "content_revision":SOURCE_REVISION,
 "head_at_build":subprocess.check_output(["git","rev-parse","HEAD"],cwd=ROOT,text=True).strip(),
 "planned_sections":sections,"main_sections":next(i for i,v in enumerate(sections) if v.startswith("附件")),"component_count":len(component_rows),"inventory_level":"framework_and_platform",
 "source_count":len(data_rows),"diagram_count":len(figures),"figures":figures,
 "screenshot_assets":[{"asset_id":f"S-{i:02}","status":"placeholder","actual_capture":False,
                       "file":f"reports/stakeholder/figures/S-{i:02}-placeholder.png"} for i in range(1,13)],
 "docx":str(path.relative_to(ROOT)),"docx_sha256":hashlib.sha256(path.read_bytes()).hexdigest(),
 "source_hashes":SOURCE_HASHES,
 "source_changed_during_build":[p for p,h in SOURCE_HASHES.items() if hashlib.sha256((ROOT/p).read_bytes()).hexdigest()!=h],
 "business_text_source":"reports/stakeholder/business-narrative.md",
 "network":False,"database_access":False,"product_tests_run":False,
 "qa_status":"awaiting_render_and_visual_inspection"}
(ROOT/"review/stakeholder-report-build-evidence.json").write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+"\n")
print(json.dumps({"docx":str(path),"planned_sections":len(sections),"diagrams":len(figures),
                  "screenshot_placeholders":12,"components":len(component_rows)},ensure_ascii=False))
